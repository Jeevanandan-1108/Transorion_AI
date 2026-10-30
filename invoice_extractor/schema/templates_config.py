"""
templates_config.py
--------------------
Single source of truth for "which template/sheet applies to which
shipment type". export/import keep their historical hardcoded filenames;
any OTHER .xlsx/.xls/.csv dropped into templates/ is discovered
automatically and registered under a --type/`type=` name of its own. See
schema.py's load_template_schema for how each of the three formats is
actually parsed into column headers.
"""
from __future__ import annotations
import re
from pathlib import Path

import openpyxl

TEMPLATES_DIR = Path(__file__).resolve().parents[2] / "templates"

_EXPLICIT_TEMPLATE_PATH: dict[str, Path] = {
    "export": TEMPLATES_DIR / "EXP_TEMPLET.xlsx",
    "import": TEMPLATES_DIR / "IMP_TEMPLET.xlsx",
}
_EXPLICIT_ITEM_SHEET: dict[str, str] = {"export": "ITEM", "import": "BOE"}

# Friendlier public-facing `?type=`/`--type` names for "export"/"import" --
# these are your ORIGINAL, hand-tuned templates (_MAPPERS/_DOMAIN_RULES in
# validate.py, _GENERIC_TO_REAL_ITEM_FIELDS in llm_extract.py, and
# detect_type.py's auto-detection are all keyed by the literal strings
# "export"/"import", independent of this file), so renaming them here
# alone would silently stop that hand-tuned logic from being found and
# fall back to the generic mapper -- the exact empty-ITEM-sheet bug this
# session spent so long diagnosing for OTHER templates. Resolved instead
# as a pure additive alias: api.py normalizes an alias straight back to its
# canonical "export"/"import" value before anything else ever sees it, so
# every hand-tuned dispatch table above keeps working completely
# unchanged, under either name.
TYPE_ALIASES: dict[str, str] = {
    "Focus export Template": "export",
    "Focus import Template": "import",
}

# A .csv template has no sheets at all -- it's one flat table. Treated as
# a single sheet under this fixed name (schema.py's load_template_schema
# must build its single dict entry under this SAME name -- there's no file
# to read it back from, so the two are kept in sync by hand).
_CSV_SHEET_NAME = "ITEMS"

_FILENAME_TO_TYPE: dict[str, str] = {
    "ExportXLSXTemplate": "Logysis export Template",
    "iCaffeExcel_Export": "iCaffe export Template",
    "Import- LOGYSIS TEMPLATE": "Logysis import Template",
    "ImportExcel (3)": "iCaffe import Template",

    "MASTERS DETAILS": "masters_details",

}


def _slugify(stem: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", stem.lower()).strip("_")
    return slug or stem.lower()


_ITEM_SHEET_NAMES = ("ITEM", "ITEMS")


def _detect_item_sheet(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _CSV_SHEET_NAME
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        for name in wb.sheet_names():
            if name.strip().upper() in _ITEM_SHEET_NAMES:
                return name
        return wb.sheet_names()[0]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            if name.strip().upper() in _ITEM_SHEET_NAMES:
                return name
        return wb.sheetnames[0]
    finally:
        wb.close()


def _discover_templates() -> tuple[dict[str, Path], dict[str, str]]:
    template_path = dict(_EXPLICIT_TEMPLATE_PATH)
    item_sheet_name = dict(_EXPLICIT_ITEM_SHEET)
    claimed = set(template_path.values())

    if not TEMPLATES_DIR.exists():
        return template_path, item_sheet_name

    paths = sorted(
        p for pattern in ("*.xlsx", "*.xls", "*.csv")
        for p in TEMPLATES_DIR.glob(pattern)
    )
    for path in paths:
        if path in claimed or path.name.startswith("~$"):
            continue
        type_name = _FILENAME_TO_TYPE.get(path.stem, _slugify(path.stem))
        template_path[type_name] = path
        item_sheet_name[type_name] = _detect_item_sheet(path)

    return template_path, item_sheet_name


TEMPLATE_PATH, ITEM_SHEET_NAME = _discover_templates()